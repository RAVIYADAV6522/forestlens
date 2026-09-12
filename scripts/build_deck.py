#!/usr/bin/env python
"""Build the ForestLens interview deck as a .pptx.

Generated rather than hand-made so the figures and every number stay tied to the
run that produced them. Regenerate the figure JPEGs first:

    python scripts/build_figures.py
    python scripts/build_deck.py --out docs/report/ForestLens.pptx

Every slide carries speaker notes: the deck is meant to be presented, and the notes
hold the answers to the questions each slide invites.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "docs" / "report" / "figures"

INK = RGBColor(0x11, 0x14, 0x18)
INK_SOFT = RGBColor(0x3B, 0x42, 0x50)
MUTED = RGBColor(0x6B, 0x74, 0x82)
RULE = RGBColor(0xD8, 0xDD, 0xD9)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
FOREST = RGBColor(0x1F, 0x6B, 0x47)
FOREST_DARK = RGBColor(0x14, 0x51, 0x2F)
WASH = RGBColor(0xEE, 0xF4, 0xF0)
AMBER = RGBColor(0x9A, 0x62, 0x06)
AMBER_WASH = RGBColor(0xFD, 0xF6, 0xE7)

BODY = "Calibri"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.72)


def textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def para(frame, text, *, size=16, bold=False, color=INK, space_after=6,
         first=False, align=PP_ALIGN.LEFT, italic=False, spacing=1.0):
    p = frame.paragraphs[0] if first else frame.add_paragraph()
    p.text = text
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = spacing
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = BODY
    return p


def rect(slide, left, top, width, height, fill, *, line=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if radius:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(slide, kicker, title, number):
    """Slide furniture: eyebrow, title, hairline rule, page number."""
    head = textbox(slide, MARGIN, Inches(0.46), W - 2 * MARGIN, Inches(1.05))
    para(head, kicker.upper(), size=11, bold=True, color=FOREST, space_after=3, first=True)
    para(head, title, size=27, bold=True, color=INK, space_after=0, spacing=0.95)

    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, MARGIN, Inches(1.62), W - MARGIN, Inches(1.62)
    )
    line.line.color.rgb = RULE
    line.line.width = Pt(1)

    num = textbox(slide, W - Inches(1.3), H - Inches(0.62), Inches(0.6), Inches(0.3))
    para(num, str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT, first=True)

    foot = textbox(slide, MARGIN, H - Inches(0.62), Inches(6), Inches(0.3))
    para(foot, "ForestLens  ·  tree-crown detection and canopy-area estimation",
         size=10, color=MUTED, first=True)


def bullets(slide, items, *, left=MARGIN, top=Inches(1.95), width=None, size=15):
    """items: list of (bold lead, rest). An empty lead renders a plain line."""
    width = width or (W - 2 * MARGIN)
    frame = textbox(slide, left, top, width, H - top - Inches(0.9))
    for i, (lead, rest) in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(11)
        p.line_spacing = 1.14
        if lead:
            run = p.add_run()
            run.text = f"{lead} "
            run.font.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = INK
            run.font.name = BODY
        run = p.add_run()
        run.text = rest
        run.font.size = Pt(size)
        run.font.color.rgb = INK_SOFT if lead else INK
        run.font.name = BODY
    return frame


def table(slide, rows, left, top, width, *, col_widths=None, size=12, highlight=None):
    n_rows, n_cols = len(rows), len(rows[0])
    height = Inches(0.34) * n_rows
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, share in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(width * share / total))

    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(0.32)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(value)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = FOREST_DARK
            elif highlight is not None and r == highlight:
                cell.fill.fore_color.rgb = WASH
            else:
                cell.fill.fore_color.rgb = PAPER
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            for run in p.runs:
                run.font.size = Pt(size)
                run.font.name = BODY
                run.font.bold = r == 0 or (highlight is not None and r == highlight)
                run.font.color.rgb = PAPER if r == 0 else INK
    return tbl


def callout(slide, text, left, top, width, *, kind="forest", size=14, height=Inches(0.95)):
    fill, bar, ink = (WASH, FOREST, INK) if kind == "forest" else (AMBER_WASH, AMBER, INK)
    rect(slide, left, top, width, height, fill)
    rect(slide, left, top, Inches(0.045), height, bar)
    frame = textbox(slide, left + Inches(0.26), top + Inches(0.16),
                    width - Inches(0.5), height - Inches(0.3))
    para(frame, text, size=size, color=ink, space_after=0, spacing=1.12, first=True)


# --------------------------------------------------------------------------- #
# Slides
# --------------------------------------------------------------------------- #

def slide_title(prs):
    slide = blank(prs)
    rect(slide, 0, 0, W, H, FOREST_DARK)
    rect(slide, 0, 0, Inches(0.13), H, FOREST)

    frame = textbox(slide, Inches(1.1), Inches(2.1), Inches(10.5), Inches(3.2))
    para(frame, "FLORA CARBON AI  ·  HIRING HACKATHON", size=12, bold=True,
         color=RGBColor(0x8F, 0xC7, 0xA9), space_after=16, first=True)
    para(frame, "ForestLens", size=54, bold=True, color=PAPER, space_after=10, spacing=0.9)
    para(frame, "Counting tree crowns and bounding canopy area in\nhigh-resolution forest imagery",
         size=21, color=RGBColor(0xD6, 0xE7, 0xDC), space_after=22, spacing=1.1)
    para(frame, "Ravi Yadav  ·  September 2026", size=15, color=RGBColor(0xA9, 0xCE, 0xBA),
         space_after=4)
    para(frame, "Live demo + source: github.com/RAVIYADAV6522/forestlens", size=13,
         color=RGBColor(0x8F, 0xC7, 0xA9), space_after=0)

    notes(slide,
        "Open with the one-sentence pitch:\n\n"
        "\"ForestLens turns high-resolution forest imagery into individual tree detections and "
        "an auditable canopy-area estimate — and it shows you where it fails.\"\n\n"
        "Then say what the deck is about: three findings that constrain what this kind of tool "
        "can honestly claim. Don't lead with the technology; lead with the findings.")


def slide_brief(prs):
    slide = blank(prs)
    chrome(slide, "The task", "The brief, and the bar it set", 2)
    bullets(slide, [
        ("Input.", "High-resolution satellite imagery of a forest. Source it yourself."),
        ("Task 1.", "Count the individual tree crowns."),
        ("Task 2.", "Estimate the canopy area those crowns cover."),
        ("Constraints.", "Solo, one weekend, any stack."),
    ], top=Inches(1.95), width=Inches(6.1))

    right = Inches(7.1)
    frame = textbox(slide, right, Inches(1.95), Inches(5.5), Inches(0.4))
    para(frame, "JUDGED ON THREE THINGS", size=11, bold=True, color=FOREST, first=True)
    table(slide, [
        ["Criterion", "What it demanded"],
        ["Does it run?", "A real application, not a concept"],
        ["Can a stranger use it?", "No developer standing beside it"],
        ["Honest about limits?", "No invented figures"],
    ], right, Inches(2.35), Inches(5.5), col_widths=[2, 3], size=12)

    callout(slide, "The organiser's stated preference: a rough tool that admits its limits beats "
                   "a polished tool that invents figures. That sentence shaped every design "
                   "decision that follows.",
            right, Inches(4.15), Inches(5.5), height=Inches(1.15))

    notes(slide,
        "Make the point that the third criterion is the hard one. Anyone can get a detector to "
        "emit a number; the difficulty is knowing which numbers you're entitled to report.\n\n"
        "If asked why canopy area is harder than counting: counting needs the crowns to be "
        "separable, area needs them to be complete. Those fail in different conditions — which "
        "is slide 8.")


def slide_rule(prs):
    slide = blank(prs)
    chrome(slide, "Design principle", "One rule drove the whole build", 3)

    callout(slide, "If the system cannot justify a number, it does not present that number.",
            MARGIN, Inches(1.9), W - 2 * MARGIN, size=19, height=Inches(0.78))

    bullets(slide, [
        ("No ground sample distance →", "no physical area. Pixel counts only, stated as such."),
        ("No segmentation masks →", "area comes from boxes, labelled a proxy everywhere, with the measured 12–22% overestimate shown."),
        ("No detector installed →", "the app refuses to count rather than inventing a number. There is deliberately no synthetic fallback."),
        ("Confidence ≠ accuracy.", "No accuracy figure is published anywhere in the project, because no labelled evaluation set exists."),
        ("Cover is an interval.", "When the two independent estimates disagree by more than 25 points, no single figure is reported at all."),
        ("No carbon estimate.", "Canopy area is geometry. Converting it needs allometry, species and field calibration I don't have."),
    ], top=Inches(2.95), size=14)

    notes(slide,
        "This is the slide to spend time on. Each line is enforced in code and covered by a "
        "test — they are not aspirations.\n\n"
        "The last one usually gets a reaction at a carbon company. Be ready to defend it: "
        "declining to convert area to carbon is a correctness property, not a missing feature. "
        "Offering a carbon number from RGB geometry alone would be the single least defensible "
        "thing this tool could do.")


def slide_data(prs):
    slide = blank(prs)
    chrome(slide, "Data", "Three scenes, self-sourced, licences verified", 4)

    table(slide, [
        ["Scene", "Source", "Native GSD", "Grid", "Extent", "Canopy"],
        ["S1a open", "Maxar/Vantor WV-02", "0.49 m", "0.305 m", "9.77 ha", "open pine/fir"],
        ["S1b dense", "Maxar/Vantor WV-02", "0.49 m", "0.305 m", "9.77 ha", "near-closed conifer"],
        ["S2 closed", "swisstopo SWISSIMAGE", "0.10 m", "0.10 m", "4.19 ha", "closed mixed"],
    ], MARGIN, Inches(1.95), W - 2 * MARGIN, col_widths=[2, 3.4, 1.7, 1.4, 1.4, 2.6], size=13)

    bullets(slide, [
        ("Reproducible, not downloaded.", "Fetch scripts read only the requested window from the provider's cloud-optimised GeoTIFF, so each crop is 1–10 MB and rebuilds exactly from a scene ID and pixel window."),
        ("Provenance is machine-readable.", "Provider, platform, date, native GSD, grid spacing, CRS, off-nadir, sun elevation, cloud, processing steps and licence — written automatically at fetch time."),
        ("Native GSD ≠ grid spacing.", "Maxar delivers a 0.305 m grid from a 0.49 m sensor. Areas from the grid are correct; detectable detail is capped by the sensor. Tracked separately."),
    ], top=Inches(3.5), size=14)

    callout(slide, "Licensing is an engineering constraint, not paperwork: Maxar Open Data is "
                   "CC BY-NC — non-commercial. Productising this needs different imagery. "
                   "SWISSIMAGE permits commercial use with attribution.",
            MARGIN, Inches(6.0), W - 2 * MARGIN, kind="amber", height=Inches(0.82))

    notes(slide,
        "Expect: \"why not NEON for the 10 cm scene?\" Answer: NEON is what DeepForest was "
        "trained on, possibly including those very tiles. It would show behaviour at native "
        "resolution while testing nothing about generalisation. SWISSIMAGE gives the same "
        "resolution over structurally different forest, which separates resolution from "
        "familiarity — and that separation is what produced the headline finding.\n\n"
        "If asked about the licence: I read the terms for the exact asset, not the programme "
        "page, and recorded them. A carbon company will care.")


def slide_pipeline(prs):
    slide = blank(prs)
    chrome(slide, "System", "Pipeline, and the arithmetic behind every area", 5)

    stages = ["Imagery\nGSD, CRS", "Scale match\nto 0.10 m/px", "Tiled detect\n800 px / 0.15",
              "Suppress dup.\nIoU 0.40", "Crown mask\noptional", "Count · area\ncover interval"]
    box_w, gap, top = Inches(1.85), Inches(0.19), Inches(2.1)
    left = MARGIN
    for i, label in enumerate(stages):
        accent = i in (1, 5)
        shape = rect(slide, left, top, box_w, Inches(0.95),
                     WASH if accent else PAPER,
                     line=FOREST if accent else RULE, radius=0.06)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.05)
        head, sub = label.split("\n")
        para(frame, head, size=13, bold=True, color=INK, space_after=1,
             align=PP_ALIGN.CENTER, first=True)
        para(frame, sub, size=10, color=MUTED, space_after=0, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            arrow = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, left + box_w, top + Inches(0.475),
                left + box_w + gap, top + Inches(0.475))
            arrow.line.color.rgb = INK
            arrow.line.width = Pt(1.25)
        left = left + box_w + gap

    bullets(slide, [
        ("Crown area.", "pixel count × GSD².  Total = sum over accepted crowns."),
        ("Canopy cover.", "the UNION of crown footprints ÷ analysed ground area — a union, so overlapping detections are not double counted."),
        ("Unknown GSD.", "physical area withheld entirely; pixel areas reported instead. No default is assumed."),
        ("One implementation.", "the UI, the CSV and the GeoJSON all call the same conversion, so the tests that pin the arithmetic pin the reported numbers."),
    ], top=Inches(3.45), size=14)

    callout(slide, "Counting and area are separate stages on purpose: the count comes from the "
                   "detector, the area from whatever crown geometry is defensible for that imagery.",
            MARGIN, Inches(6.05), W - 2 * MARGIN, height=Inches(0.72))

    notes(slide,
        "Walk left to right, then give the arithmetic. Two details worth volunteering:\n\n"
        "1. Boxes found on the resampled copy are mapped back into source-raster pixels, so "
        "annotation, area and GeoJSON all share one coordinate frame.\n"
        "2. Cover uses the union, not the sum. I found the sum double-counting overlaps while "
        "building the cover bracket — the sum is still reported, but as crown area, not cover.")


def slide_finding_one(prs):
    slide = blank(prs)
    chrome(slide, "Finding 1 of 3", "The detector fails on satellite imagery at native scale", 6)

    if (FIGURES / "fig_native.jpg").exists():
        slide.shapes.add_picture(str(FIGURES / "fig_native.jpg"), MARGIN, Inches(2.0),
                                 height=Inches(2.75))
        slide.shapes.add_picture(str(FIGURES / "fig_resampled.jpg"), Inches(3.55), Inches(2.0),
                                 height=Inches(2.75))
    cap = textbox(slide, MARGIN, Inches(4.85), Inches(5.6), Inches(0.5))
    para(cap, "Same scene, same settings. Left: native 0.305 m, 57 crowns.  "
              "Right: resampled 3.05×, 695 crowns.", size=11, color=MUTED, first=True)

    right = Inches(6.75)
    table(slide, [
        ["Scale", "Eff. GSD", "Trees", "/ha", "Mean crown"],
        ["1×", "0.305 m", "57", "5.8", "179 m² (15.1 m)"],
        ["2×", "0.153 m", "392", "40.1", "49.0 m² (7.9 m)"],
        ["3.05×", "0.100 m", "695", "71.2", "21.3 m² (5.2 m)"],
        ["4×", "0.076 m", "836", "85.6", "12.1 m² (3.9 m)"],
    ], right, Inches(2.0), Inches(5.85), col_widths=[1.1, 1.4, 1, 1, 2.2], size=12, highlight=3)

    bullets(slide, [
        ("", "At native scale: 5.8 trees/ha with 15 m “crowns”. Real crowns here are 4–8 m. The model was trained on ~0.10 m imagery."),
        ("", "So the detection input is resampled to the model's training GSD — the same crop then gives 71.2 trees/ha with 5.2 m crowns."),
    ], left=right, top=Inches(3.85), width=Inches(5.85), size=14)

    callout(slide, "3.05× is anchored to the model's documented training GSD — not picked. The "
                   "count never plateaus, so choosing the nicest number would be tuning to a "
                   "preferred answer. And resampling adds no information.",
            right, Inches(5.35), Inches(5.85), kind="amber", height=Inches(1.35))

    notes(slide,
        "The trap on this slide is that 57 → 695 looks like an improvement. It is not a result; "
        "it is evidence that the input scale was wrong. Say that before they ask.\n\n"
        "\"Isn't upsampling cheating?\" — It adds no information and the app says so on screen. "
        "It presents crowns at the pixel size the model expects. What would be cheating is "
        "picking 4× because 836 looked better; since the count rises monotonically and never "
        "plateaus, the only defensible anchor is the documented training resolution.\n\n"
        "Whether 695 is right is slide 9.")


def slide_finding_two(prs):
    slide = blank(prs)
    chrome(slide, "Finding 2 of 3", "Canopy closure, not resolution, is the binding constraint", 7)

    if (FIGURES / "fig_closed.jpg").exists():
        slide.shapes.add_picture(str(FIGURES / "fig_closed.jpg"), MARGIN, Inches(2.0),
                                 height=Inches(3.5))
    cap = textbox(slide, MARGIN, Inches(5.6), Inches(3.6), Inches(0.6))
    para(cap, "S2 at native 0.10 m — the detector's own training resolution, no resampling.",
         size=11, color=MUTED, first=True)

    right = Inches(4.9)
    table(slide, [
        ["Scene", "GSD", "Trees/ha", "Cover (crown union)", "Visible cover"],
        ["S1a open", "0.49 m", "71.2", "17.0%", "consistent"],
        ["S1b dense", "0.49 m", "72.2", "15.1%", "~70–90%"],
        ["S2 closed", "0.10 m", "37.7", "21.0%", "~95–100%"],
    ], right, Inches(2.0), Inches(7.7), col_widths=[1.8, 1.1, 1.4, 2.2, 1.8], size=12, highlight=3)

    bullets(slide, [
        ("The finest scene performed worst.", "At exactly the training resolution it recovers roughly one crown in five."),
        ("Counts saturate.", "71.2 vs 72.2 trees/ha across two visibly different stand densities — the detector stops separating crowns once they interlock."),
        ("Two causes, confounded.", "Closure removes the separable apexes the detector needs; and DeepForest's training data is largely open North American conifer, not closed European mixed forest."),
    ], left=right, top=Inches(3.5), width=Inches(7.7), size=14)

    callout(slide, "This inverts the usual procurement instinct. Buying finer imagery does not "
                   "fix individual-tree analysis under closed canopy — assess stand structure "
                   "before commissioning imagery.",
            right, Inches(5.9), Inches(7.7), height=Inches(0.85))

    notes(slide,
        "This is the strongest slide in the deck and the one to lead with if time is short.\n\n"
        "I sourced the Swiss scene specifically to test this, and deliberately avoided NEON so "
        "that resolution and familiarity were separable. The result was the opposite of what I "
        "expected going in.\n\n"
        "If asked what would fix it: fine-tuning on closed-canopy imagery. No threshold setting "
        "reaches an 80% miss rate — I tried.")


def slide_finding_three(prs):
    slide = blank(prs)
    chrome(slide, "Finding 3 of 3", "Summed crown boxes are not a canopy-cover measurement", 8)

    bullets(slide, [
        ("The failure.", "Detection fires on separable crown apexes, so the union of its boxes cannot tile interlocking canopy. On the closed-canopy scene that reports 21% where cover is visibly near-complete — wrong by a factor of four."),
        ("The response.", "Report cover as two independent estimates and refuse to average them."),
    ], top=Inches(1.95), width=Inches(6.0), size=15)

    table(slide, [
        ["Scene", "Lower: crown union", "Upper: green vegetation", "Width"],
        ["S1a open", "17.0%", "92.6%", "76 pts"],
        ["S1b dense", "15.1%", "96.3%", "81 pts"],
        ["S2 closed", "21.0%", "93.9%", "73 pts"],
    ], MARGIN, Inches(3.95), Inches(6.0), col_widths=[1.6, 2, 2.3, 1.1], size=12)

    right = Inches(7.15)
    frame = textbox(slide, right, Inches(1.95), Inches(5.45), Inches(2.6))
    para(frame, "WHY EACH BOUND IS WRONG", size=11, bold=True, color=FOREST,
         space_after=8, first=True)
    para(frame, "Lower bound — union of detected crowns.", size=14, bold=True, space_after=2)
    para(frame, "Too low: every crown the detector missed contributes nothing. Under closed "
                "canopy, far too low.", size=13, color=INK_SOFT, space_after=10, spacing=1.1)
    para(frame, "Upper bound — pixels where 2G − R − B > 0.", size=14, bold=True, space_after=2)
    para(frame, "Too high: grass, shrubs and crops are green too. Deeply shadowed canopy can "
                "drop below the threshold and be missed.", size=13, color=INK_SOFT,
         space_after=0, spacing=1.1)

    callout(slide, "Above 25 points apart, the app states the range and refuses a single figure — "
                   "the midpoint of two estimates that disagree this strongly means nothing. "
                   "Honest conclusion: this pipeline counts relatively separable trees and cannot "
                   "measure canopy cover to a useful precision.",
            right, Inches(4.7), Inches(5.45), kind="amber", height=Inches(1.55))

    notes(slide,
        "Expect: \"a 75-point range isn't very useful.\" Agree — and that is the finding. A "
        "useless honest answer beats a useful false one, and it tells you which instrument you "
        "actually need.\n\n"
        "The right instrument for cover is a canopy-versus-ground segmentation model. Crown "
        "detection answers \"how many trees\"; it is the wrong tool for \"how much cover\". "
        "That's my first recommendation for next steps.")


def slide_validation(prs):
    slide = blank(prs)
    chrome(slide, "Validation", "Counted by eye, before any detections were shown", 9)

    table(slide, [
        ["Crop", "Scene", "Area", "Reference", "Detected", "Difference"],
        ["openA", "S1a, 3.05×", "0.610 ha", "48 ± 10", "57", "+19%"],
        ["openB", "S1a, 3.05×", "0.610 ha", "36 ± 8", "46", "+28%"],
        ["chA", "S2, native", "0.262 ha", "48 ± 14", "9", "−81%"],
    ], MARGIN, Inches(1.95), W - 2 * MARGIN,
        col_widths=[1.4, 2.2, 1.6, 1.8, 1.6, 1.8], size=13, highlight=3)

    bullets(slide, [
        ("Method, stated up front.", "One count per distinguishable crown unit bounded by shadow or texture discontinuity; unresolvable clumps counted by apparent apex; crops quartered and counted cell by cell."),
        ("Open canopy runs 19–28% high.", "At or just beyond the counting uncertainty. Likely clump splitting and shadow detections."),
        ("Closed canopy recovers ~1 in 5.", "Not a borderline result, and no threshold setting changes its character."),
        ("The uncertainty is itself a finding.", "At 0.49 m native GSD even careful manual counting cannot separate clumps reliably — so don't over-read the open-canopy agreement."),
    ], top=Inches(3.5), size=14)

    callout(slide, "These reference counts were made by the AI assistant that built the system — "
                   "by eye, not by a human expert or field survey. They are independent of the "
                   "detector but they are not ground truth, and both headline figures inherit "
                   "that weakness. Stated wherever the numbers appear.",
            MARGIN, Inches(6.05), W - 2 * MARGIN, kind="amber", height=Inches(0.85))

    notes(slide,
        "Do not skip the amber box. Volunteering the weakness before being asked is the whole "
        "point of the project.\n\n"
        "If you have since counted the crops yourself, replace these numbers with yours and say "
        "so — that converts the weakest part of the submission into a number you can defend as "
        "your own.\n\n"
        "\"How accurate is it?\" — I don't publish an accuracy figure. There's no labelled "
        "evaluation set for these scenes, and eyeballing an overlay doesn't give precision or "
        "recall. What I have is this table, with its provenance stated.")


def slide_rejected(prs):
    slide = blank(prs)
    chrome(slide, "Rigour", "Three things I tried, measured, and threw away", 10)

    frame = textbox(slide, MARGIN, Inches(1.95), Inches(6.05), Inches(4.6))
    para(frame, "OTSU THRESHOLDING — REJECTED", size=11, bold=True, color=FOREST,
         space_after=5, first=True)
    para(frame, "It assumes a bimodal histogram with two real classes. On a uniformly vegetated "
                "image it splits within the vegetation distribution — sunlit against shaded "
                "crown — and claimed 62% cover on a ~100%-closed scene. Replaced with a fixed, "
                "stated threshold; Otsu is still reported as a diagnostic.",
         size=13, color=INK_SOFT, space_after=16, spacing=1.12)

    para(frame, "TEXTURE FILTERING — REJECTED", size=11, bold=True, color=FOREST, space_after=5)
    para(frame, "The vegetation index reads turbid lake water as 94.3% vegetation. Canopy is "
                "rough and water is smooth, so local σ looks like the fix: water 4.8, closed "
                "canopy 20.0 — but the open Okanagan stand is 6.3. Any threshold removing the "
                "lake also deletes a real forest scene. The app flags the condition instead of "
                "shipping a fix that works on one image.",
         size=13, color=INK_SOFT, space_after=0, spacing=1.12)

    right = Inches(7.2)
    frame = textbox(slide, right, Inches(1.95), Inches(5.4), Inches(1.2))
    para(frame, "A LIBRARY DOCSTRING — CHECKED, NOT TRUSTED", size=11, bold=True,
         color=FOREST, space_after=5, first=True)
    para(frame, "The detector documents its image argument as BGR. The pipeline passes RGB. "
                "Rather than assume, both were tested on the library's own in-distribution crop:",
         size=13, color=INK_SOFT, space_after=0, spacing=1.12)

    table(slide, [
        ["Channel order", "Trees", "Mean conf.", "Conf > 0.5"],
        ["RGB (what we pass)", "55", "0.535", "35"],
        ["BGR (per the docs)", "29", "0.389", "5"],
    ], right, Inches(3.35), Inches(5.4), col_widths=[2.4, 1, 1.3, 1.3], size=12, highlight=1)

    callout(slide, "RGB is correct; the docstring is stale. Had it gone the other way, every "
                   "number in this project would have come from transposed colour channels.",
            right, Inches(4.7), Inches(5.4), kind="amber", height=Inches(0.85))

    negatives = textbox(slide, right, Inches(5.75), Inches(5.4), Inches(1.1))
    para(negatives, "NEGATIVE CONTROLS", size=11, bold=True, color=FOREST, space_after=5, first=True)
    para(negatives, "Open water: 2 detections over 9.77 ha, 0.0% crown union — effectively clean. "
                    "Urban core: detections land on real street trees, roofs almost entirely "
                    "clear, a handful of genuine rooftop false positives.",
         size=12, color=INK_SOFT, space_after=0, spacing=1.1)

    notes(slide,
        "This slide exists to show method rather than results. Interviewers can tell the "
        "difference between someone who got a pipeline working and someone who interrogated it.\n\n"
        "The channel-order check is the strongest single anecdote in the deck: a stale docstring "
        "would have silently invalidated every count, and the only reason it didn't is that I "
        "tested the claim instead of trusting it.\n\n"
        "The texture rejection is the one that shows judgement: I had a fix that made the demo "
        "scene look right and declined to ship it because it would break elsewhere.")


def slide_engineering(prs):
    slide = blank(prs)
    chrome(slide, "Engineering", "Reproducibility, tested where it matters", 11)

    bullets(slide, [
        ("80 tests, no model weights needed.", "Area arithmetic, GSD precedence, resampling and box remapping, union-vs-sum cover, duplicate suppression, edge flagging, mask acceptance, export shapes — so the measurement logic is verifiable without a GPU."),
        ("13 mechanical QA checks.", "Run against the acceptance criteria, including a grep for accuracy claims that should not exist."),
        ("Every run exports its own metadata.", "Image identity, GSD and its origin, model versions, all thresholds, resampling scale, timings, peak memory. Any count is reproducible from its own export."),
        ("Imagery rebuilds from a scene ID.", "Window-addressed COG reads, not a folder of downloads."),
    ], top=Inches(1.95), width=Inches(6.15), size=14)

    right = Inches(7.25)
    frame = textbox(slide, right, Inches(1.95), Inches(5.35), Inches(0.4))
    para(frame, "MEMORY — MEASURED, NOT ASSUMED", size=11, bold=True, color=FOREST, first=True)
    table(slide, [
        ["Inference path", "Peak RSS", "Detections"],
        ["In-memory tiling", "~1270 MB", "695"],
        ["Windowed (default)", "~1100 MB", "695, identical"],
    ], right, Inches(2.4), Inches(5.35), col_widths=[2.4, 1.5, 1.8], size=12, highlight=2)

    frame = textbox(slide, right, Inches(3.6), Inches(5.35), Inches(2.4))
    para(frame, "Windowed inference reads one tile at a time from a temporary raster. Verified "
                "bit-identical: same count, same area, same hash of all box coordinates.",
         size=13, color=INK_SOFT, space_after=10, spacing=1.12, first=True)
    para(frame, "Two interventions tested and discarded: disabling gradients changed nothing "
                "(the library already does it), and a 400 px tile did not cut memory while "
                "changing the result — 1379 trees against 695. Tile size is a far stronger "
                "lever on the count than on memory.",
         size=13, color=INK_SOFT, space_after=0, spacing=1.12)

    callout(slide, "Run-to-run variance is ±150 MB and the floor is ~1050 MB even on the small "
                   "scene, so the cost is the model's forward pass — not the imagery.",
            right, Inches(5.95), Inches(5.35), height=Inches(0.8))

    notes(slide,
        "Use this slide if the interviewer is technical. Two things to emphasise:\n\n"
        "1. The tests cover the arithmetic, not the model. That's deliberate — I can't test "
        "DeepForest's weights, but I can guarantee that pixel count × GSD² is right and that an "
        "unknown GSD withholds area.\n"
        "2. I found and fixed a real defect through this: the area formula existed in three "
        "places, and the one the tests covered was the one the app never called. Consolidating "
        "it meant the tests finally pinned the shipped numbers.")


def slide_next(prs):
    slide = blank(prs)
    chrome(slide, "Limits and next steps", "What I would not claim, and what I'd build next", 12)

    frame = textbox(slide, MARGIN, Inches(1.95), Inches(6.05), Inches(4.4))
    para(frame, "WHAT I WOULD NOT CLAIM", size=11, bold=True, color=AMBER,
         space_after=8, first=True)
    for lead, rest in [
        ("No accuracy figure.", "No labelled evaluation set exists for these scenes."),
        ("Reference counts aren't ground truth.", "Made by eye by the assistant that built the tool."),
        ("Three scenes, two forest types.", "Closure is confounded with structural domain shift."),
        ("The upper cover bound is approximate.", "It over-counts green water and under-counts shadowed canopy."),
        ("Generalisation is untested.", "Including — honestly — on Indian forests, which look structurally closer to the failure case."),
    ]:
        p = frame.add_paragraph()
        p.space_after = Pt(9)
        p.line_spacing = 1.1
        run = p.add_run(); run.text = f"{lead} "
        run.font.bold = True; run.font.size = Pt(13.5); run.font.color.rgb = INK; run.font.name = BODY
        run = p.add_run(); run.text = rest
        run.font.size = Pt(13.5); run.font.color.rgb = INK_SOFT; run.font.name = BODY

    right = Inches(7.15)
    frame = textbox(slide, right, Inches(1.95), Inches(5.45), Inches(4.4))
    para(frame, "WHAT I WOULD BUILD NEXT", size=11, bold=True, color=FOREST,
         space_after=8, first=True)
    for i, (lead, rest) in enumerate([
        ("Expert or field reference counts.", "Converts every “plausible” here into a measurement."),
        ("Fine-tune on closed canopy.", "Targets the dominant error — the 80% miss rate no threshold reaches."),
        ("A canopy-vs-ground segmentation model.", "Replaces the 75-point cover interval with an actual measurement."),
        ("Height from lidar or photogrammetry.", "Separates trees from grass, which RGB greenness cannot."),
    ], start=1):
        p = frame.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = 1.1
        run = p.add_run(); run.text = f"{i}.  {lead} "
        run.font.bold = True; run.font.size = Pt(13.5); run.font.color.rgb = INK; run.font.name = BODY
        run = p.add_run(); run.text = rest
        run.font.size = Pt(13.5); run.font.color.rgb = INK_SOFT; run.font.name = BODY

    callout(slide, "“Why should you trust the result?  Because I can show you every detected "
                   "crown, where the resolution came from, the arithmetic behind the area, and "
                   "the scene where it fails worst — none of it hidden behind a single "
                   "confident number.”",
            MARGIN, Inches(6.05), W - 2 * MARGIN, size=15, height=Inches(0.9))

    notes(slide,
        "Close on the quote — it is the thesis of the whole project in one sentence.\n\n"
        "The ordering of next steps is deliberate and worth saying out loud: reference counts "
        "first, because without them every figure in the deck is only plausible. Then "
        "fine-tuning, because it attacks the dominant error. The segmentation model third, "
        "because it fixes the wrong-instrument problem for cover.\n\n"
        "If asked \"what would you do differently?\" — I'd source the closed-canopy scene on day "
        "one. It produced the most important finding and I found it late.")


def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for slide_fn in (
        slide_title, slide_brief, slide_rule, slide_data, slide_pipeline,
        slide_finding_one, slide_finding_two, slide_finding_three,
        slide_validation, slide_rejected, slide_engineering, slide_next,
    ):
        slide_fn(prs)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=ROOT / "docs/report/ForestLens-deck.pptx")
    args = parser.parse_args()
    path = build(args.out)
    print(f"wrote {path} ({path.stat().st_size / 1e6:.2f} MB)")
